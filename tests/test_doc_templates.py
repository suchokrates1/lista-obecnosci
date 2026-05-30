import os
import pytest
from docx import Document

from doc_generator import generuj_liste_obecnosci, generuj_raport_miesieczny

class Dummy:
    pass


def test_missing_liste_template():
    if os.path.exists('szablon.docx'):
        os.remove('szablon.docx')
    with pytest.raises(RuntimeError) as exc:
        generuj_liste_obecnosci('2024-01-01', '1', [], 'T', None)
    assert 'szablon.docx' in str(exc.value)


def test_missing_report_template(tmp_path):
    path = tmp_path / 'missing.docx'
    dummy = Dummy()
    dummy.imie = 'A'
    dummy.nazwisko = 'B'
    dummy.numer_umowy = '1'
    dummy.podpis_filename = 'sig.png'
    with pytest.raises(RuntimeError) as exc:
        generuj_raport_miesieczny(dummy, [], str(path), 'static', 1, 2024)
    assert str(path) in str(exc.value)


def test_attendance_template_fills_assistant_name(tmp_path, monkeypatch):
    template = tmp_path / 'szablon.docx'
    doc = Document()
    doc.add_paragraph('Lista obecności')
    doc.add_paragraph('Data zajęć:    Czas trwania zajęć:')
    attendees = doc.add_table(rows=2, cols=1)
    attendees.cell(0, 0).text = 'Uczestnik'
    footer = doc.add_table(rows=2, cols=4)
    footer.cell(0, 0).text = 'Imię i nazwisko trenera'
    footer.cell(0, 1).text = 'Podpis trenera'
    footer.cell(0, 2).text = 'Imię i nazwisko asystenta'
    footer.cell(0, 3).text = 'Podpis asystenta'
    doc.save(template)

    monkeypatch.chdir(tmp_path)
    sig = tmp_path / 'assistant.png'
    sig.write_bytes(
        bytes.fromhex(
            '89504E470D0A1A0A0000000D4948445200000001000000010802000000907753DE'
            '0000000C49444154789C63F8CFC0000003010100C9FE92EF0000000049454E44AE426082'
        )
    )

    generated = generuj_liste_obecnosci(
        '2026-05-11',
        '1,5',
        ['Anna'],
        'Trener Testowy',
        None,
        assistant_name='Asystent Testowy',
        assistant_signature_path=str(sig),
    )

    assert generated.tables[1].cell(1, 0).text == 'Trener Testowy'
    assert generated.tables[1].cell(1, 2).text == 'Asystent Testowy'
    assert 'pic:pic' in generated.tables[1].cell(1, 3)._tc.xml
